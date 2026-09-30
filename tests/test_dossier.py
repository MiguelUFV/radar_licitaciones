"""El dossier en Word dice lo mismo que el PDF.

El riesgo que vigilan estas pruebas no es que el Word salga feo: es que diga **otra cosa**.
Había dos .docx escritos a mano el 24-09-2026 que se quedaron viejos en cuanto se midió, y el
arreglo fue generarlo del mismo material que el PDF. Si alguien vuelve a separar los dos
caminos, esto se pone en rojo.
"""

from __future__ import annotations

import pytest

pytest.importorskip("reportlab", reason="el dossier se genera con --with reportlab")
pytest.importorskip("docx", reason="el dossier se genera con --with python-docx")

from docx import Document  # noqa: E402

from docs import dossier  # noqa: E402


def leer(ruta) -> str:
    documento = Document(str(ruta))
    partes = [p.text for p in documento.paragraphs]
    for tabla in documento.tables:
        partes += [celda.text for fila in tabla.rows for celda in fila.cells]
    return "\n".join(partes)


def test_el_marcado_del_pdf_no_se_cuela_en_el_word(tmp_path):
    """Un `<b>` o un `&nbsp;` sin traducir se leería tal cual en Word."""
    historia = [
        dossier.p("Lo <b>medido</b> son 4,7&nbsp;páginas", "cuerpo"),
        dossier.p("Una línea<br/>y otra", "nota"),
    ]
    texto = leer(dossier.a_docx(historia, tmp_path / "d.docx"))

    assert "<b>" not in texto and "&nbsp;" not in texto and "<br/>" not in texto
    assert "Lo medido son 4,7 páginas" in texto
    assert "Una línea\ny otra" in texto


def test_las_cifras_de_una_tabla_llegan_al_word(tmp_path):
    """Una tabla del PDF son Paragraph dentro de celdas; si se leyeran mal, saldría vacía."""
    historia = [dossier.tabla([["Métrica", "Valor"], ["Recall del agente", "69,2 %"]], [100, 100])]
    texto = leer(dossier.a_docx(historia, tmp_path / "d.docx"))

    assert "Recall del agente" in texto
    assert "69,2 %" in texto


def test_los_titulos_son_titulos_de_word(tmp_path):
    """Sin esto el Word sale como un bloque de texto sin índice ni navegación."""
    historia = [dossier.p("4. Qué se midió", "seccion"), dossier.p("Un apartado", "sub")]
    documento = Document(str(dossier.a_docx(historia, tmp_path / "d.docx")))

    niveles = {p.text: p.style.name for p in documento.paragraphs if p.text}
    assert niveles["4. Qué se midió"] == "Heading 1"
    assert niveles["Un apartado"] == "Heading 2"


def test_la_regla_horizontal_no_sale_como_una_tabla_vacia(tmp_path):
    """`regla()` es una tabla de una celda vacía; copiada tal cual dejaría cuadrículas sueltas."""
    documento = Document(str(dossier.a_docx([dossier.regla()], tmp_path / "d.docx")))

    assert documento.tables == []


def test_el_resultado_del_estudio_sale_de_la_base_y_no_escrito_a_mano():
    """El fallo que motivó esto: el dossier decía 74,0 % y «19 de 20» semanas después de que
    fueran 69,2 % y 16 de 16, porque estaban escritos a mano en cuatro sitios."""
    hoy = {
        "resultado": {
            "baseline_b": {
                "diferencia": -0.1923,
                "ic": (-0.3269, -0.0769),
                "contratos": 52,
                "recall_agente": 0.6923,
                "recall_baseline": 0.8846,
                "veredicto": "refutada",
            }
        }
    }
    fila = dossier.comparacion(hoy, "baseline_b", "Frente al filtro hecho a medida")

    assert fila[1:4] == ["69,2 %", "88,5 %", "<b>−19,2 puntos</b>"]
    assert "tesis refutada" in fila[4]
    assert dossier.perdidos(hoy) == "16"


def test_sin_base_el_dossier_lo_dice_en_lugar_de_inventarse_el_resultado():
    """Un documento que trata de que nada sea inventado no puede inventarse su propia tabla."""
    fila = dossier.comparacion({}, "baseline_b", "Frente al filtro hecho a medida")

    assert fila[1:4] == ["—", "—", "—"]
    assert "no respondió" in fila[4]
    assert dossier.perdidos({}) == "casi todos los"
